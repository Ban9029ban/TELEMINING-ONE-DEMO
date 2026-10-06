from __future__ import annotations

from datetime import datetime, timedelta
from io import BytesIO
from typing import Any, Dict, Iterable, List

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

GANTT_SLOT_LABELS = [
    "20 a 21",
    "21 a 22",
    "22 a 23",
    "23 a 00",
    "00 a 01",
    "01 a 02",
    "02 a 03",
    "03 a 04",
    "04 a 05",
    "05 a 06",
    "06 a 07",
    "07 a 08",
    "08 a 09",
]

_THIN = Side(style="thin", color="CBD5E1")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
_HEADER_FILL = PatternFill("solid", fgColor="0F172A")
_SECTION_FILL = PatternFill("solid", fgColor="1E3A8A")
_POSTURA_FILL = PatternFill("solid", fgColor="1D4ED8")
_PROGRAM_FILL = PatternFill("solid", fgColor="2563EB")
_PROGRAM_RESOURCE_FILL = PatternFill("solid", fgColor="DBEAFE")
_REAL_FILL = PatternFill("solid", fgColor="111827")
_REAL_RESOURCE_FILL = PatternFill("solid", fgColor="E2E8F0")
_META_FILL = PatternFill("solid", fgColor="F8FAFC")
_WHITE_FONT = Font(color="FFFFFF", bold=True)
_DARK_FONT = Font(color="0F172A", bold=True)


def _parse_dt(value: Any) -> datetime | None:
    if not value:
        return None
    if isinstance(value, datetime):
        return value.replace(tzinfo=None)
    try:
        return datetime.fromisoformat(str(value).replace("Z", "+00:00")).replace(tzinfo=None)
    except ValueError:
        return None


def _normalize_operation_date(value: Any) -> datetime:
    parsed = _parse_dt(value)
    if parsed:
        return parsed.replace(hour=0, minute=0, second=0, microsecond=0)
    return datetime.utcnow().replace(hour=0, minute=0, second=0, microsecond=0)


def _normalize_task_dt(value: Any, operation_date: Any) -> datetime | None:
    parsed = _parse_dt(value)
    if not parsed:
        return None

    base_date = _normalize_operation_date(operation_date)
    if parsed.date() == base_date.date() and parsed.hour < 12:
        parsed = parsed + timedelta(days=1)
    return parsed


def _slot_bounds(operation_date: Any, slot_label: str) -> tuple[datetime, datetime]:
    base_date = _normalize_operation_date(operation_date)
    start_label, end_label = [part.strip() for part in slot_label.split("a")]
    start_hour = int(start_label)
    end_hour = int(end_label)

    start = base_date.replace(hour=start_hour, minute=0, second=0, microsecond=0)
    if start_hour < 12:
        start += timedelta(days=1)

    end = base_date.replace(hour=end_hour, minute=0, second=0, microsecond=0)
    if end_hour <= start_hour or end_hour < 12:
        end += timedelta(days=1)

    return start, end


def _task_window(task: Dict[str, Any], start_key: str, end_key: str) -> tuple[datetime | None, datetime | None]:
    operation_date = task.get("date") or task.get("operationDate")
    start = _normalize_task_dt(task.get(start_key), operation_date)
    end = _normalize_task_dt(task.get(end_key), operation_date)

    if start and end and end <= start:
        end += timedelta(days=1)

    return start, end


def _overlaps(
    task: Dict[str, Any],
    slot_start: datetime,
    slot_end: datetime,
    start_key: str,
    end_key: str,
) -> bool:
    start, end = _task_window(task, start_key, end_key)
    if not start or not end:
        return False
    return start < slot_end and end > slot_start


def _compose_resource(task: Dict[str, Any], prefix: str) -> str:
    equipment = str(task.get(f"equipment{prefix}") or "").strip()
    crew = str(task.get(f"crew{prefix}") or "").strip()
    return " / ".join(part for part in (equipment, crew) if part)


def _dedupe_lines(values: Iterable[str]) -> str:
    seen = set()
    ordered = []
    for raw in values:
        value = str(raw or "").strip()
        if not value or value in seen:
            continue
        seen.add(value)
        ordered.append(value)
    return "\n".join(ordered)


def _compose_observations(tasks: List[Dict[str, Any]]) -> str:
    seen = set()
    ordered = []
    for task in tasks:
        note = str(task.get("observation") or "").strip()
        if not note or note in seen:
            continue
        seen.add(note)
        ordered.append(f"- {note}")
    return "\n".join(ordered)


def _task_sort_key(task: Dict[str, Any]) -> tuple[datetime, str, str]:
    operation_date = task.get("date") or task.get("operationDate")
    planned_start, _ = _task_window(task, "plannedStart", "plannedEnd")
    real_start, _ = _task_window(task, "realStart", "realEnd")
    base = planned_start or real_start or _normalize_operation_date(operation_date)
    return (base, str(task.get("nivel") or ""), str(task.get("postura") or ""))


def _build_hours_from_tasks(tasks: List[Dict[str, Any]]) -> Dict[str, List[Dict[str, Any]]]:
    crew_rows: Dict[str, Dict[str, Any]] = {}
    equipment_rows: Dict[str, Dict[str, Any]] = {}

    for task in tasks:
        planned_minutes = int(task.get("plannedMinutes") or 0)
        real_minutes = int(task.get("realMinutes") or 0)
        planned_hours = round(planned_minutes / 60, 2)
        real_hours = round(real_minutes / 60, 2)

        crew_plan = str(task.get("crewPlanned") or "").strip()
        crew_real = str(task.get("crewReal") or "").strip()
        equipment_plan = str(task.get("equipmentPlanned") or "").strip()
        equipment_real = str(task.get("equipmentReal") or "").strip()

        if crew_plan:
            current = crew_rows.setdefault(
                crew_plan,
                {
                    "label": crew_plan,
                    "plannedHours": 0,
                    "realHours": 0,
                    "targetHours": 0,
                    "plannedTasks": 0,
                    "realTasks": 0,
                },
            )
            current["plannedHours"] += planned_hours
            current["targetHours"] += planned_hours
            current["plannedTasks"] += 1

        if crew_real:
            current = crew_rows.setdefault(
                crew_real,
                {
                    "label": crew_real,
                    "plannedHours": 0,
                    "realHours": 0,
                    "targetHours": 0,
                    "plannedTasks": 0,
                    "realTasks": 0,
                },
            )
            current["realHours"] += real_hours
            current["realTasks"] += 1

        if equipment_plan:
            current = equipment_rows.setdefault(
                equipment_plan,
                {
                    "label": equipment_plan,
                    "plannedHours": 0,
                    "realHours": 0,
                    "targetHours": 0,
                    "plannedTasks": 0,
                    "realTasks": 0,
                },
            )
            current["plannedHours"] += planned_hours
            current["targetHours"] += planned_hours
            current["plannedTasks"] += 1

        if equipment_real:
            current = equipment_rows.setdefault(
                equipment_real,
                {
                    "label": equipment_real,
                    "plannedHours": 0,
                    "realHours": 0,
                    "targetHours": 0,
                    "plannedTasks": 0,
                    "realTasks": 0,
                },
            )
            current["realHours"] += real_hours
            current["realTasks"] += 1

    crew = sorted(crew_rows.values(), key=lambda item: item["plannedHours"], reverse=True)
    equipment = sorted(equipment_rows.values(), key=lambda item: item["plannedHours"], reverse=True)
    return {"crew": crew, "equipment": equipment}


def build_gantt_rows(tasks: List[Dict[str, Any]], slot_labels: List[str] | None = None) -> List[Dict[str, Any]]:
    slot_labels = slot_labels or GANTT_SLOT_LABELS
    grouped: Dict[str, Dict[str, Any]] = {}

    for task in sorted(tasks, key=_task_sort_key):
        key = str(task.get("waypointId") or f"{task.get('nivel') or '-'}::{task.get('postura') or '-'}")
        current = grouped.get(key)
        if not current:
            current = {
                "waypointId": task.get("waypointId") or key,
                "nivel": task.get("nivel") or "",
                "postura": task.get("postura") or "",
                "mapCode": task.get("mapCode") or "",
                "turno": task.get("turno") or "",
                "operationDate": task.get("date") or task.get("operationDate"),
                "tasks": [],
            }
            grouped[key] = current
        current["tasks"].append(task)

    rows: List[Dict[str, Any]] = []
    for group in grouped.values():
        group_tasks = sorted(group["tasks"], key=_task_sort_key)
        operation_date = group["operationDate"]
        slots = []

        for slot_label in slot_labels:
            slot_start, slot_end = _slot_bounds(operation_date, slot_label)
            planned_tasks = [
                task
                for task in group_tasks
                if _overlaps(task, slot_start, slot_end, "plannedStart", "plannedEnd")
            ]
            real_tasks = [
                task
                for task in group_tasks
                if _overlaps(task, slot_start, slot_end, "realStart", "realEnd")
            ]

            slots.append(
                {
                    "slot": slot_label,
                    "program": _dedupe_lines(task.get("activity") or task.get("type") or "-" for task in planned_tasks),
                    "resourcePlan": _dedupe_lines(_compose_resource(task, "Planned") for task in planned_tasks),
                    "real": _dedupe_lines(task.get("activity") or task.get("type") or "-" for task in real_tasks),
                    "resourceReal": _dedupe_lines(_compose_resource(task, "Real") for task in real_tasks),
                }
            )

        rows.append(
            {
                "waypointId": group["waypointId"],
                "nivel": group["nivel"],
                "postura": group["postura"],
                "mapCode": group["mapCode"],
                "turno": group["turno"],
                "operationDate": operation_date,
                "observations": _compose_observations(group_tasks),
                "tasks": group_tasks,
                "slots": slots,
            }
        )

    rows.sort(
        key=lambda item: (
            _task_sort_key(item["tasks"][0])[0] if item.get("tasks") else _normalize_operation_date(item.get("operationDate")),
            item.get("nivel") or "",
            item.get("postura") or "",
        )
    )
    return rows


def _format_percent(value: Any) -> str:
    return f"{round(float(value or 0) * 100)}%"


def _format_date(value: Any) -> str:
    parsed = _parse_dt(value)
    return parsed.strftime("%d-%m-%Y") if parsed else "-"


def _format_time(value: Any, operation_date: Any = None) -> str:
    parsed = _normalize_task_dt(value, operation_date) if operation_date else _parse_dt(value)
    return parsed.strftime("%H:%M") if parsed else "-"


def _map_label(meta_maps: List[Dict[str, Any]], map_code: str) -> str:
    for item in meta_maps or []:
        if str(item.get("code")) == str(map_code):
            return str(item.get("label") or map_code)
    return map_code or "Mapa"


def _fill_cell(cell, *, fill=None, font=None, align=None, border=True):
    if fill is not None:
        cell.fill = fill
    if font is not None:
        cell.font = font
    if align is not None:
        cell.alignment = align
    if border:
        cell.border = _BORDER


def _write_gantt_sheet(
    ws,
    *,
    title: str,
    rows: List[Dict[str, Any]],
    compliance: Any,
    meta_maps: List[Dict[str, Any]],
    operation_date: Any = None,
    turn_label: str | None = None,
    display_title: str = "Control y Seguimiento Programa",
):
    last_column = 5 + len(GANTT_SLOT_LABELS)
    last_column_letter = get_column_letter(last_column)
    resolved_operation_date = operation_date or (rows[0]["operationDate"] if rows else None)
    resolved_turn_label = turn_label or (rows[0]["turno"] if rows else "-")
    map_values = sorted({str(row.get("mapCode") or "") for row in rows if row.get("mapCode")})
    map_label = _map_label(meta_maps, map_values[0]) if len(map_values) == 1 else "Multiples mapas"

    ws.title = title
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "F6"

    ws["A1"] = "Fecha"
    ws["B1"] = _format_date(resolved_operation_date)
    ws["D1"] = "Cumplimiento Programa"
    ws["E1"] = _format_percent(compliance)
    ws["A2"] = "Turno"
    ws["B2"] = resolved_turn_label or "-"
    ws["D2"] = "Mapa"
    ws["E2"] = map_label

    ws.merge_cells(f"A4:{last_column_letter}4")
    ws["A4"] = display_title

    for ref in ("A1", "D1", "A2", "D2"):
        _fill_cell(ws[ref], fill=_SECTION_FILL, font=_WHITE_FONT, align=Alignment(horizontal="center"))
    for ref in ("B1", "E1", "B2", "E2"):
        _fill_cell(ws[ref], fill=_META_FILL, font=_DARK_FONT, align=Alignment(horizontal="center"))

    _fill_cell(
        ws["A4"],
        fill=_HEADER_FILL,
        font=Font(color="FFFFFF", bold=True, size=14),
        align=Alignment(horizontal="center"),
    )

    headers = ["#", "Nivel", "Postura", "Plan", "Observaciones", *GANTT_SLOT_LABELS]
    for column_index, header in enumerate(headers, start=1):
        cell = ws.cell(row=5, column=column_index, value=header)
        _fill_cell(cell, fill=_HEADER_FILL, font=_WHITE_FONT, align=Alignment(horizontal="center", vertical="center"))

    widths = [6, 10, 24, 16, 42] + [13] * len(GANTT_SLOT_LABELS)
    for index, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(index)].width = width

    plan_styles = {
        "program": (_PROGRAM_FILL, Font(color="FFFFFF", bold=True)),
        "resourcePlan": (_PROGRAM_RESOURCE_FILL, Font(color="0F172A", bold=False)),
        "real": (_REAL_FILL, Font(color="FFFFFF", bold=True)),
        "resourceReal": (_REAL_RESOURCE_FILL, Font(color="0F172A", bold=False)),
    }
    line_labels = {
        "program": "Programa",
        "resourcePlan": "Recurso Prog",
        "real": "Real",
        "resourceReal": "Recurso Real",
    }

    current_row = 6
    for row_index, row in enumerate(rows, start=1):
        slot_map = {item["slot"]: item for item in row.get("slots", [])}
        group_start = current_row
        group_end = current_row + 3

        ws.merge_cells(start_row=group_start, start_column=1, end_row=group_end, end_column=1)
        ws.merge_cells(start_row=group_start, start_column=2, end_row=group_end, end_column=2)
        ws.merge_cells(start_row=group_start, start_column=3, end_row=group_end, end_column=3)
        ws.merge_cells(start_row=group_start, start_column=5, end_row=group_end, end_column=5)

        ws.cell(row=group_start, column=1, value=row_index)
        ws.cell(row=group_start, column=2, value=row.get("nivel") or "")
        ws.cell(row=group_start, column=3, value=row.get("postura") or "")
        ws.cell(row=group_start, column=5, value=row.get("observations") or "")

        _fill_cell(ws.cell(row=group_start, column=1), fill=_META_FILL, font=_DARK_FONT, align=Alignment(horizontal="center", vertical="center"))
        _fill_cell(ws.cell(row=group_start, column=2), fill=_META_FILL, font=_DARK_FONT, align=Alignment(horizontal="center", vertical="center"))
        _fill_cell(ws.cell(row=group_start, column=3), fill=_POSTURA_FILL, font=_WHITE_FONT, align=Alignment(horizontal="center", vertical="center", wrap_text=True))
        _fill_cell(ws.cell(row=group_start, column=5), fill=_META_FILL, font=Font(color="0F172A"), align=Alignment(vertical="top", wrap_text=True))

        for line_offset, line_key in enumerate(("program", "resourcePlan", "real", "resourceReal")):
            line_row = group_start + line_offset
            label_cell = ws.cell(row=line_row, column=4, value=line_labels[line_key])
            fill, font = plan_styles[line_key]
            _fill_cell(label_cell, fill=fill, font=font, align=Alignment(horizontal="center", vertical="center"))

            for slot_index, slot_label in enumerate(GANTT_SLOT_LABELS, start=6):
                value = slot_map.get(slot_label, {}).get(line_key, "")
                cell = ws.cell(row=line_row, column=slot_index, value=value or "")
                _fill_cell(cell, fill=fill, font=font, align=Alignment(horizontal="center", vertical="center", wrap_text=True))

            ws.row_dimensions[line_row].height = 26

        current_row = group_end + 1


def _control_columns() -> List[tuple[str, str]]:
    return [
        ("Fecha", "date"),
        ("Turno", "turno"),
        ("Nivel", "nivel"),
        ("Postura", "postura"),
        ("Tipo", "type"),
        ("Actividad", "activity"),
        ("Prioridad", "priority"),
        ("Hora Inicio Prog", "plannedStart"),
        ("Hora Término Prog", "plannedEnd"),
        ("Equipo Prog", "equipmentPlanned"),
        ("Cuadrilla Prog", "crewPlanned"),
        ("Avance (%)", "progress"),
        ("Hora Inicio Real", "realStart"),
        ("Hora Término Real", "realEnd"),
        ("Equipo Real", "equipmentReal"),
        ("Cuadrilla Real", "crewReal"),
        ("Categoría CNC Externa", "cncExternalCategory"),
        ("Subcategoría CNC Externa", "cncExternalSubcategory"),
        ("Detalle CNC Externa", "cncExternalDetail"),
        ("Hora Inicio CNC Externa", "cncExternalStart"),
        ("Hora Término CNC Externa", "cncExternalEnd"),
        ("Categoría CNC Interna", "cncInternalCategory"),
        ("Subcategoría CNC Interna", "cncInternalSubcategory"),
        ("Detalle CNC Interna", "cncInternalDetail"),
        ("Hora Inicio CNC Interna", "cncInternalStart"),
        ("Hora Término CNC Interna", "cncInternalEnd"),
        ("Cargo 1", "role1"),
        ("Cargo 2", "role2"),
        ("Cargo 3", "role3"),
        ("Observaciones", "observation"),
        ("HIP", "hip"),
        ("HTP", "htp"),
        ("HIR", "hir"),
        ("HTR", "htr"),
        ("HIE", "hie"),
        ("HTE", "hte"),
        ("HII", "hii"),
        ("HTI", "hti"),
        ("Pórtico Tracking", "trackingPorticoName"),
        ("Personas Vinculadas", "linkedPeople"),
        ("Vehículos Vinculados", "linkedVehicles"),
    ]


def _control_value(task: Dict[str, Any], key: str) -> Any:
    operation_date = task.get("date") or task.get("operationDate")
    if key == "date":
        return _format_date(task.get("date"))
    if key == "progress":
        return _format_percent(task.get("progress"))
    if key.startswith("role"):
        index = int(key[-1]) - 1
        roles = list(task.get("roles") or [])
        return roles[index] if index < len(roles) else ""
    if key == "linkedPeople":
        return ", ".join(item.get("label") or item.get("name") or "" for item in task.get("linkedPeople") or [])
    if key == "linkedVehicles":
        return ", ".join(item.get("label") or item.get("name") or "" for item in task.get("linkedVehicles") or [])
    if key in {"plannedStart", "plannedEnd", "realStart", "realEnd", "cncExternalStart", "cncExternalEnd", "cncInternalStart", "cncInternalEnd", "hip", "htp", "hir", "htr", "hie", "hte", "hii", "hti"}:
        return _format_time(task.get(key), operation_date)
    return task.get(key) or ""


def _write_control_sheet(ws, *, title: str, tasks: List[Dict[str, Any]]):
    ws.title = title
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "A2"

    columns = _control_columns()
    for column_index, (label, _key) in enumerate(columns, start=1):
        cell = ws.cell(row=1, column=column_index, value=label)
        _fill_cell(cell, fill=_HEADER_FILL, font=_WHITE_FONT, align=Alignment(horizontal="center", vertical="center", wrap_text=True))
        ws.column_dimensions[get_column_letter(column_index)].width = 18 if column_index != 30 else 32

    for row_index, task in enumerate(sorted(tasks, key=_task_sort_key), start=2):
        for column_index, (_label, key) in enumerate(columns, start=1):
            cell = ws.cell(row=row_index, column=column_index, value=_control_value(task, key))
            _fill_cell(
                cell,
                fill=_META_FILL if row_index % 2 == 0 else PatternFill("solid", fgColor="FFFFFF"),
                font=Font(color="0F172A"),
                align=Alignment(vertical="top", wrap_text=True),
            )


def _write_hours_sheet(ws, *, title: str, hours_payload: Dict[str, Any]):
    ws.title = title
    ws.sheet_view.showGridLines = False
    ws["A1"] = "Cuadrillas"
    ws["F1"] = "Equipos"
    _fill_cell(ws["A1"], fill=_SECTION_FILL, font=_WHITE_FONT, align=Alignment(horizontal="center"))
    _fill_cell(ws["F1"], fill=_SECTION_FILL, font=_WHITE_FONT, align=Alignment(horizontal="center"))

    crew_headers = ["Etiqueta", "Horas Plan", "Horas Real", "Meta", "Tareas Plan"]
    equipment_headers = ["Etiqueta", "Horas Plan", "Horas Real", "Meta", "Tareas Plan"]
    for idx, header in enumerate(crew_headers, start=1):
        _fill_cell(ws.cell(row=2, column=idx, value=header), fill=_HEADER_FILL, font=_WHITE_FONT, align=Alignment(horizontal="center"))
    for idx, header in enumerate(equipment_headers, start=6):
        _fill_cell(ws.cell(row=2, column=idx, value=header), fill=_HEADER_FILL, font=_WHITE_FONT, align=Alignment(horizontal="center"))

    crew = list(hours_payload.get("crew") or [])
    equipment = list(hours_payload.get("equipment") or [])
    max_rows = max(len(crew), len(equipment), 1)
    for offset in range(max_rows):
        crew_item = crew[offset] if offset < len(crew) else {}
        equipment_item = equipment[offset] if offset < len(equipment) else {}
        row_number = 3 + offset

        crew_values = [
            crew_item.get("label") or "",
            crew_item.get("plannedHours") or 0,
            crew_item.get("realHours") or 0,
            crew_item.get("targetHours") or 0,
            crew_item.get("plannedTasks") or 0,
        ]
        equipment_values = [
            equipment_item.get("label") or "",
            equipment_item.get("plannedHours") or 0,
            equipment_item.get("realHours") or 0,
            equipment_item.get("targetHours") or 0,
            equipment_item.get("plannedTasks") or 0,
        ]

        for idx, value in enumerate(crew_values, start=1):
            _fill_cell(
                ws.cell(row=row_number, column=idx, value=value),
                fill=_META_FILL if row_number % 2 else PatternFill("solid", fgColor="FFFFFF"),
                font=Font(color="0F172A"),
                align=Alignment(horizontal="center"),
            )
        for idx, value in enumerate(equipment_values, start=6):
            _fill_cell(
                ws.cell(row=row_number, column=idx, value=value),
                fill=_META_FILL if row_number % 2 else PatternFill("solid", fgColor="FFFFFF"),
                font=Font(color="0F172A"),
                align=Alignment(horizontal="center"),
            )

    for column in range(1, 11):
        ws.column_dimensions[get_column_letter(column)].width = 18


def build_production_workbook_bytes(
    payload: Dict[str, Any],
    *,
    mineria_tasks: List[Dict[str, Any]] | None = None,
    oocc_tasks: List[Dict[str, Any]] | None = None,
    focus: str | None = None,
) -> BytesIO:
    schedule = payload.get("schedule") or {}
    meta_maps = list((payload.get("meta") or {}).get("maps") or [])
    mineria_tasks = list(mineria_tasks if mineria_tasks is not None else schedule.get("controlMineria") or [])
    oocc_tasks = list(oocc_tasks if oocc_tasks is not None else schedule.get("controlOOCC") or [])

    workbook = Workbook()
    workbook.remove(workbook.active)

    if focus in (None, "mineria"):
        mineria_rows = build_gantt_rows(mineria_tasks)
        _write_gantt_sheet(
            workbook.create_sheet("Gantt Minería"),
            title="Gantt Minería",
            rows=mineria_rows,
            compliance=(payload.get("dashboard") or {}).get("cumplimientoMineria"),
            meta_maps=meta_maps,
            operation_date=(payload.get("meta") or {}).get("operationDate"),
            display_title="Control y Seguimiento Programa",
        )
        _write_control_sheet(
            workbook.create_sheet("Control Minería"),
            title="Control Minería",
            tasks=mineria_tasks,
        )
        _write_hours_sheet(
            workbook.create_sheet("Horas Minería"),
            title="Horas Minería",
            hours_payload=_build_hours_from_tasks(mineria_tasks),
        )

    if focus in (None, "oocc"):
        oocc_rows = build_gantt_rows(oocc_tasks)
        _write_gantt_sheet(
            workbook.create_sheet("Gantt OO.CC"),
            title="Gantt OO.CC",
            rows=oocc_rows,
            compliance=(payload.get("dashboard") or {}).get("cumplimientoOOCC"),
            meta_maps=meta_maps,
            operation_date=(payload.get("meta") or {}).get("operationDate"),
            display_title="Control y Seguimiento Programa",
        )
        _write_control_sheet(
            workbook.create_sheet("Control OO.CC"),
            title="Control OO.CC",
            tasks=oocc_tasks,
        )

    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output
