import base64
import logging
from datetime import date, datetime, timedelta
from io import BytesIO
from typing import Optional, Dict, Any

from fastapi import APIRouter, Body, Depends, Query
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.utils import get_column_letter
from sqlalchemy import text
from sqlalchemy.ext.asyncio import AsyncSession

from app.db.db import get_db

router = APIRouter(prefix="/control/graphics", tags=["graphics"])
logger = logging.getLogger(__name__)


def _sanitize_sheet_name(name: str) -> str:
    # Excel limits sheet names to 31 chars and forbids: : \ / ? * [ ]
    bad = set(":\\/?*[]")
    safe = "".join("_" if c in bad else c for c in (name or ""))
    safe = (safe or "Hoja")[:31]
    return safe


def _best_fit_columns(ws, max_width: int = 60):
    """Ajuste simple de ancho de columnas según contenido."""
    dims = {}
    for row in ws.iter_rows(values_only=True):
        for idx, value in enumerate(row, start=1):
            if value is None:
                continue
            s = str(value)
            dims[idx] = max(dims.get(idx, 0), min(max_width, len(s) + 2))

    for idx, width in dims.items():
        ws.column_dimensions[get_column_letter(idx)].width = width


def _build_datetime_range(
    start_date: Optional[date],
    end_date: Optional[date],
) -> tuple[datetime, datetime]:
    """
    Devuelve un rango [start_dt, end_dt) en base a fechas.
    - Si no se recibe nada -> día actual completo.
    - Si sólo viene start_date -> ese día completo.
    - Si start_date > end_date, se intercambian.
    """
    if start_date is None and end_date is None:
        d = date.today()
        start_date = d
        end_date = d
    elif start_date is None and end_date is not None:
        start_date = end_date
    elif start_date is not None and end_date is None:
        end_date = start_date

    if start_date > end_date:
        start_date, end_date = end_date, start_date

    start_dt = datetime.combine(start_date, datetime.min.time())
    end_dt = datetime.combine(end_date + timedelta(days=1), datetime.min.time())
    return start_dt, end_dt


async def _table_exists(db: AsyncSession, table_name: str) -> bool:
    result = await db.execute(
        text(
            """
            SELECT EXISTS (
                SELECT 1
                FROM information_schema.tables
                WHERE table_schema = 'public'
                  AND table_name = :table_name
            )
            """
        ),
        {"table_name": table_name},
    )
    return bool(result.scalar())


async def _get_table_columns(db: AsyncSession, table_name: str) -> set[str]:
    result = await db.execute(
        text(
            """
            SELECT column_name
            FROM information_schema.columns
            WHERE table_schema = 'public'
              AND table_name = :table_name
            """
        ),
        {"table_name": table_name},
    )
    return {str(name) for name in result.scalars().all()}


def _resolve_seismic_polygon_status_expr(columns: set[str]) -> str:
    direct_candidates = (
        "status",
        "state",
        "polygon_status",
        "status_name",
        "status_value",
        "semaforo",
        "semaphore_status",
        "traffic_light",
    )
    for candidate in direct_candidates:
        if candidate in columns:
            return candidate

    severity_col = next((candidate for candidate in ("severity", "alert_severity") if candidate in columns), None)
    source_state_col = next((candidate for candidate in ("source_state", "sync_state") if candidate in columns), None)

    if not severity_col and not source_state_col:
        return "'Unknown'"

    severity_expr = f"COALESCE({severity_col}::text, '')" if severity_col else "''"
    source_expr = f"LOWER(COALESCE({source_state_col}::text, ''))" if source_state_col else "''"

    return f"""
    CASE
      WHEN {source_expr} = 'offline' THEN 'Offline'
      WHEN {source_expr} = 'unknown' THEN 'Unknown'
      WHEN {severity_expr} IN ('Critica', 'Critical') THEN 'Critical'
      WHEN {severity_expr} IN ('Alta', 'High', 'Warning') THEN 'Warning'
      WHEN {severity_expr} IN ('Media', 'Medium', 'Offline') THEN 'Offline'
      WHEN {severity_expr} IN ('Baja', 'Low', 'Normal') THEN 'Normal'
      WHEN {severity_expr} = 'Unknown' THEN 'Unknown'
      ELSE 'Normal'
    END
    """.strip()


def _resolve_seismic_row_id_expr(columns: set[str]) -> str:
    for candidate in ("id", "created_at", "polygon_name", "site_name", "plan_name", "observed_at"):
        if candidate in columns:
            return candidate
    return "observed_at"


def _build_seismic_order_clause(row_id_expr: str, direction: str = "ASC") -> str:
    parts = [f"observed_at {direction}"]
    if row_id_expr and row_id_expr != "observed_at":
        parts.append(f"{row_id_expr} {direction}")
    return ", ".join(parts)


def _empty_seismic_dashboard_response(
    bucket: str = "day",
    *,
    degraded: bool = False,
    error: str | None = None,
) -> dict[str, Any]:
    return {
        "kpis": {
            "events_in_range": 0,
            "critical_changes": 0,
            "open_alerts": 0,
            "polygons_monitored": 0,
            "sync_runs": 0,
            "evacuation_hours_total": 0,
            "latest_data_at": None,
        },
        "sites": [],
        "plans": [],
        "interval_bucket": bucket,
        "events_by_interval": [],
        "site_panorama": [],
        "alerts_by_severity": [],
        "top_polygons": [],
        "polling_correlation": [],
        "source_state": [],
        "activity_by_plan": [],
        "freshness_by_plan": [],
        "degraded": degraded,
        "error": error,
    }



@router.get("/entities")
async def list_entities(
    entity_type: str | None = Query(default=None, description="person | vehicle (opcional)"),
    db: AsyncSession = Depends(get_db),
):
    """Lista entidades disponibles para filtros (personas y vehículos).

    Devuelve:
    - Si entity_type es None -> { persons: [...], vehicles: [...] }
    - Si entity_type=person  -> { persons: [...] }
    - Si entity_type=vehicle -> { vehicles: [...] }

    Cada item: { id: "<id>", name: "<display_name>" }
    """

    out = {}

    if entity_type is None or entity_type == "person":
        q = await db.execute(text("""
            SELECT p.id::text AS id, COALESCE(p.full_name, p.rut::text, p.id::text) AS name
            FROM person p
            ORDER BY COALESCE(p.full_name, p.rut::text, p.id::text);
        """))
        out["persons"] = [dict(r) for r in q.mappings().all()]

    if entity_type is None or entity_type == "vehicle":
        # Preferimos un display_name "bonito" (ej: "aljibe KK-LL-66") desde tag_assignment.
        # Fallback: patente.
        q = await db.execute(text("""
            SELECT
              v.id::text AS id,
              COALESCE(
                CASE WHEN ta.display_name IS NOT NULL AND POSITION(' ' IN ta.display_name) > 0 THEN ta.display_name END,
                CASE
                  WHEN v.status IS NOT NULL
                   AND v.status NOT IN ('online','offline')
                   AND v.patente IS NOT NULL
                  THEN (v.status || ' ' || v.patente)
                  ELSE v.patente
                END,
                v.id::text
              ) AS name,
              v.patente AS patente
            FROM vehicles v
            LEFT JOIN LATERAL (
              SELECT a.display_name
              FROM tag_assignment a
              WHERE a.active = TRUE
                AND a.entity_type = 'vehicle'
                AND a.entity_id = v.id
                AND a.display_name IS NOT NULL
              ORDER BY a.updated_at DESC NULLS LAST, a.created_at DESC
              LIMIT 1
            ) ta ON TRUE
            ORDER BY COALESCE(ta.display_name, v.patente, v.id::text);
        """))
        out["vehicles"] = [dict(r) for r in q.mappings().all()]

    return out


@router.post("/export-excel")
async def export_excel_with_images(payload: dict = Body(...)):
    """Genera un XLSX con data + imágenes (PNG base64) por hoja.

    Espera un JSON con:
      - filename: str (opcional)
      - filters: dict (opcional)
      - sheets: [
          {
            "name": "Permanencias",
            "rows": [ {..}, {..} ],
            "image": "data:image/png;base64,..." (opcional)
          }, ...
        ]
    """

    filename = (payload.get("filename") or "tmone_export.xlsx").strip() or "tmone_export.xlsx"
    filters = payload.get("filters") or {}
    sheets = payload.get("sheets") or []

    wb = Workbook()
    # eliminar hoja por defecto
    if wb.active is not None:
        wb.remove(wb.active)

    # ---- Hoja filtros ----
    ws_filters = wb.create_sheet(_sanitize_sheet_name("Filtros"))
    ws_filters.cell(row=1, column=1, value="campo")
    ws_filters.cell(row=1, column=2, value="valor")
    r = 2
    for k, v in (filters or {}).items():
        ws_filters.cell(row=r, column=1, value=str(k))
        ws_filters.cell(row=r, column=2, value="" if v is None else str(v))
        r += 1
    _best_fit_columns(ws_filters)

    # ---- Hojas data + imágenes ----
    for s in sheets:
        name = _sanitize_sheet_name(str(s.get("name") or "Hoja"))
        ws = wb.create_sheet(name)

        # título
        ws.cell(row=1, column=1, value=str(s.get("name") or name))

        # imagen (si viene)
        img_data_url = s.get("image")
        start_row = 3
        if isinstance(img_data_url, str) and img_data_url.startswith("data:image"):
            try:
                b64 = img_data_url.split(",", 1)[1]
                img_bytes = base64.b64decode(b64)
                img = XLImage(BytesIO(img_bytes))
                img.anchor = "A2"
                ws.add_image(img)
                # Dejamos un espacio para que no se mezcle con la tabla
                start_row = 25
            except Exception:
                # Si falla la imagen, igual exportamos la data
                start_row = 3

        rows = s.get("rows") or []
        if isinstance(rows, list) and rows:
            # encabezados según primera fila (dict)
            if isinstance(rows[0], dict):
                headers = list(rows[0].keys())
            else:
                headers = ["value"]

            # header
            for c, h in enumerate(headers, start=1):
                ws.cell(row=start_row, column=c, value=h)

            # data
            cur = start_row + 1
            for item in rows:
                if isinstance(item, dict):
                    for c, h in enumerate(headers, start=1):
                        ws.cell(row=cur, column=c, value=item.get(h))
                else:
                    ws.cell(row=cur, column=1, value=item)
                cur += 1

            _best_fit_columns(ws)
        else:
            ws.cell(row=start_row, column=1, value="(sin datos)")

    bio = BytesIO()
    wb.save(bio)
    bio.seek(0)

    headers = {
        "Content-Disposition": f'attachment; filename="{filename}"'
    }
    return StreamingResponse(
        bio,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers=headers,
    )

# -------------------------------------------------------------
# PERSONAS POR NIVEL (NP, UCL, NTI, OFICINA, etc.)
# -------------------------------------------------------------
@router.get("/people-per-level")
async def people_per_level(
    db: AsyncSession = Depends(get_db),
    start_date: Optional[date] = Query(None, description="Fecha inicial (YYYY-MM-DD)"),
    end_date: Optional[date] = Query(None, description="Fecha final (YYYY-MM-DD)"),
    nivel: Optional[str] = Query(None, description="Filtrar por nivel (NP, UCL, NTI, etc.)"),
    portico_id: Optional[int] = Query(None, description="Filtrar por pórtico (access_portico.id)"),
) -> Dict[str, int]:
    """
    Devuelve un dict {nivel: total_personas} en el rango de fechas indicado,
    contando personas distintas vistas en cada nivel.
    """
    start_dt, end_dt = _build_datetime_range(start_date, end_date)

    filters = ["ae.ts_edge >= :start_dt", "ae.ts_edge < :end_dt", "ta.entity_type = 'person'", "ta.active = TRUE"]
    params: Dict[str, Any] = {"start_dt": start_dt, "end_dt": end_dt}

    if nivel:
        filters.append("COALESCE(mp.nivel, 'SIN NIVEL') = :nivel")
        params["nivel"] = nivel

    if portico_id:
        filters.append("(ae.extra ->> 'portico')::bigint = :portico_id")
        params["portico_id"] = portico_id

    where_clause = " AND ".join(filters)

    sql = text(f"""
        SELECT
            COALESCE(mp.nivel, 'SIN NIVEL') AS nivel,
            COUNT(DISTINCT ta.entity_id)     AS total
        FROM access_event ae
        JOIN tag_assignment ta
          ON ta.tag_id = ae.tag_id
        LEFT JOIN access_portico ap
          ON ap.id = (ae.extra ->> 'portico')::bigint
        LEFT JOIN map_points mp
          ON mp.id = ap.map_point_id
        WHERE {where_clause}
        GROUP BY COALESCE(mp.nivel, 'SIN NIVEL')
        ORDER BY COALESCE(mp.nivel, 'SIN NIVEL');
    """)

    q = await db.execute(sql, params)
    rows = q.fetchall()

    # Formato que espera el front: { "NP": 10, "UCL": 5, ... }
    return {row[0]: int(row[1]) for row in rows}


# -------------------------------------------------------------
# VEHÍCULOS POR NIVEL Y TIPO (usamos status como "tipo" básico)
# -------------------------------------------------------------
@router.get("/vehicles-per-level")
async def vehicles_per_level(
    db: AsyncSession = Depends(get_db),
    start_date: Optional[date] = Query(None),
    end_date: Optional[date] = Query(None),
    nivel: Optional[str] = Query(None),
    portico_id: Optional[int] = Query(None),
) -> Dict[str, Dict[str, int]]:
    """
    Devuelve un dict anidado:
      {
        "NP":  { "online": 12, "offline": 3 },
        "UCL": { "online": 4 },
        ...
      }
    donde la clave externa es el nivel (NP/UCL/NTI/SIN NIVEL) y la interna el "tipo"
    (usamos el campo vehicles.status como categoría).
    """
    start_dt, end_dt = _build_datetime_range(start_date, end_date)

    filters = ["ae.ts_edge >= :start_dt", "ae.ts_edge < :end_dt", "ta.entity_type = 'vehicle'", "ta.active = TRUE"]
    params: Dict[str, Any] = {"start_dt": start_dt, "end_dt": end_dt}

    if nivel:
        filters.append("COALESCE(mp.nivel, 'SIN NIVEL') = :nivel")
        params["nivel"] = nivel

    if portico_id:
        filters.append("(ae.extra ->> 'portico')::bigint = :portico_id")
        params["portico_id"] = portico_id

    where_clause = " AND ".join(filters)

    sql = text(f"""
        SELECT
            COALESCE(mp.nivel, 'SIN NIVEL') AS nivel,
            COALESCE(v.status, 'sin_tipo')  AS tipo,
            COUNT(DISTINCT ta.entity_id)    AS total
        FROM access_event ae
        JOIN tag_assignment ta
          ON ta.tag_id = ae.tag_id
        JOIN vehicles v
          ON v.id = ta.entity_id
        LEFT JOIN access_portico ap
          ON ap.id = (ae.extra ->> 'portico')::bigint
        LEFT JOIN map_points mp
          ON mp.id = ap.map_point_id
        WHERE {where_clause}
        GROUP BY COALESCE(mp.nivel, 'SIN NIVEL'), COALESCE(v.status, 'sin_tipo')
        ORDER BY 1, 2;
    """)

    q = await db.execute(sql, params)
    rows = q.fetchall()

    # Reorganizar a dict anidado por nivel -> tipo
    result: Dict[str, Dict[str, int]] = {}
    for nivel_val, tipo, total in rows:
        nivel_key = nivel_val or "SIN NIVEL"
        if nivel_key not in result:
            result[nivel_key] = {}
        result[nivel_key][tipo] = int(total)

    return result


# -------------------------------------------------------------
# VEHICLE STATUS (donut "Vehículos conectados")
# -------------------------------------------------------------
@router.get("/vehicle-status")
async def vehicle_status(
    db: AsyncSession = Depends(get_db),
    minutes_window: int = Query(5, ge=1, le=1440),
) -> Dict[str, int]:
    """
    Retorna algo como:
      { "Conectados": 15, "Sin datos en ventana": 8 }
    usando la última detección de cada vehículo en access_event
    dentro de la ventana indicada (en minutos).
    """
    now = datetime.utcnow()
    window_start = now - timedelta(minutes=minutes_window)

    # Última detección por vehículo
    sql_last_seen = text("""
        WITH last_seen AS (
            SELECT
                ta.entity_id AS vehicle_id,
                MAX(ae.ts_edge) AS ts_last
            FROM access_event ae
            JOIN tag_assignment ta
              ON ta.tag_id = ae.tag_id
            WHERE ta.entity_type = 'vehicle'
              AND ta.active = TRUE
            GROUP BY ta.entity_id
        )
        SELECT
            SUM(CASE WHEN ts_last >= :window_start THEN 1 ELSE 0 END) AS conectados,
            SUM(CASE WHEN ts_last <  :window_start THEN 1 ELSE 0 END) AS sin_datos
        FROM last_seen;
    """)

    q = await db.execute(sql_last_seen, {"window_start": window_start})
    row = q.fetchone()
    conectados = int(row[0] or 0)
    sin_datos = int(row[1] or 0)

    return {
        "Conectados": conectados,
        "Sin datos en ventana": sin_datos,
    }


# -------------------------------------------------------------
# TMOne STATUS (usa view_zenit_kpis creada en 00_init.sql)
# -------------------------------------------------------------
@router.get("/zenit-status")
async def zenit_status(db: AsyncSession = Depends(get_db)):
    """
    Lee directamente la vista view_zenit_kpis, que ya existe en 00_init.sql:
      - URs operativas
      - URs con alerta
      - Vehículos operativos
    El front la puede mostrar como prefiera.
    """
    q = await db.execute(text("SELECT categoria, cantidad FROM view_zenit_kpis;"))
    rows = q.fetchall()

    return {str(r[0]): int(r[1]) for r in rows}


# -------------------------------------------------------------
# GRAFIC DATA PARA SensorChart (clima por pórtico)
# -------------------------------------------------------------
@router.get("/grafic-data")
async def grafic_data(
    db: AsyncSession = Depends(get_db),
    start_date: Optional[date] = Query(None),
    end_date: Optional[date] = Query(None),
    portico_id: Optional[int] = Query(None),
) -> Dict[str, Any]:
    """
    Endpoint usado por SensorChart.jsx.
    Devuelve:
      {
        "sensor_data": {
          "UR1": { "timestamps": [...], "values": [...] },
          "UR2": { ... }
        },
        "critical_values": {
          "Valor critico superior": X,
          "Valor critico inferior": Y
        }
      }
    Para simplicidad tomamos:
      - Lecturas desde access_portico_climate_reading
      - Umbrales desde access_portico_climate_cfg (si existen),
        si no, se usan 18 / 26 grados por defecto.
    """
    start_dt, end_dt = _build_datetime_range(start_date, end_date)

    filters = ["r.ts >= :start_dt", "r.ts < :end_dt"]
    params: Dict[str, Any] = {"start_dt": start_dt, "end_dt": end_dt}

    if portico_id:
        filters.append("r.portico_id = :portico_id")
        params["portico_id"] = portico_id

    where_clause = " AND ".join(filters)

    # Lecturas climáticas por pórtico (nombre del pórtico)
    sql_readings = text(f"""
        SELECT
            ap.name          AS portico_name,
            r.ts             AS ts,
            r.value::float   AS value
        FROM access_portico_climate_reading r
        JOIN access_portico ap
          ON ap.id = r.portico_id
        WHERE {where_clause}
        ORDER BY ap.name, r.ts;
    """)

    q = await db.execute(sql_readings, params)
    rows = q.fetchall()

    sensor_data: Dict[str, Dict[str, list]] = {}
    for portico_name, ts, value in rows:
        if portico_name not in sensor_data:
            sensor_data[portico_name] = {"timestamps": [], "values": []}
        sensor_data[portico_name]["timestamps"].append(ts.isoformat())
        sensor_data[portico_name]["values"].append(float(value))

    # Valores críticos: tomamos el primer cfg cualquiera como referencia,
    # o default 18/26 si no existe config.
    sql_cfg = text("""
        SELECT green_max, yellow_max
        FROM access_portico_climate_cfg
        WHERE active = TRUE
        LIMIT 1;
    """)

    q_cfg = await db.execute(sql_cfg)
    cfg = q_cfg.fetchone()
    if cfg:
        lower = float(cfg[0])  # podría ser umbral inferior / verde
        upper = float(cfg[1])  # umbral amarillo
    else:
        lower = 18.0
        upper = 26.0

    critical_values = {
        "Valor critico inferior": lower,
        "Valor critico superior": upper,
    }

    return {
        "sensor_data": sensor_data,
        "critical_values": critical_values,
    }




@router.get("/vehicle-laps")
async def vehicle_laps(
    db: AsyncSession = Depends(get_db),
    start_date: date | None = Query(None, description="YYYY-MM-DD"),
    end_date: date | None = Query(None, description="YYYY-MM-DD"),
    nivel: str | None = Query(None, description="Nivel (ej: NP)"),
    portico_id: int | None = Query(None, description="ID del pórtico"),
    vehicle_id: int | None = Query(None, description="ID del vehículo (opcional)"),
):
    """Vueltas por vehículo (REAL): cuenta lecturas de access_event asociadas a vehículos.

    - start_date / end_date: rango de fechas (end_date incluye todo el día).
    - nivel: filtra por access_portico.nivel (ej: NP, UCL, ...).
    - portico_id: filtra por access_portico.id.
    - vehicle_id: filtra por un vehículo específico (vehicles.id).
    """

    start_dt, end_dt = _build_datetime_range(start_date, end_date)

    sql = text(
        """
        SELECT
          ta.entity_id::text                 AS vehicle_id,
          COALESCE(
            MAX(CASE WHEN ta.display_name IS NOT NULL AND POSITION(' ' IN ta.display_name) > 0 THEN ta.display_name END),
            MAX(
              CASE
                WHEN v.status IS NOT NULL
                 AND v.status NOT IN ('online','offline')
                 AND v.patente IS NOT NULL
                THEN (v.status || ' ' || v.patente)
                ELSE v.patente
              END
            ),
            ta.entity_id::text
          ) AS vehiculo,
          COUNT(*)::int                      AS vueltas
        FROM access_event ae
        JOIN tag_assignment ta
          ON ta.tag_id = ae.tag_id
         AND ta.active = TRUE
         AND ta.entity_type = 'vehicle'
        LEFT JOIN vehicles v
          ON v.id = ta.entity_id
        JOIN access_portico ap
          ON ap.id = CAST(ae.extra ->> 'portico' AS bigint)
         AND ap.active = TRUE
        WHERE
          COALESCE(ae.ts_edge, ae.ts_ingested) >= :start_dt
          AND COALESCE(ae.ts_edge, ae.ts_ingested) <  :end_dt
          -- NOTE: asyncpg no puede inferir el tipo de un parámetro NULL en prepared statements.
          -- Forzamos el tipo para evitar: "could not determine data type of parameter".
          AND (CAST(:nivel AS text) IS NULL OR ap.nivel::text = CAST(:nivel AS text))
          AND (CAST(:portico_id AS bigint) IS NULL OR ap.id = CAST(:portico_id AS bigint))
          AND (CAST(:vehicle_id AS bigint) IS NULL OR ta.entity_id = CAST(:vehicle_id AS bigint))
        GROUP BY ta.entity_id
        ORDER BY vueltas DESC
        LIMIT 2000;
        """
    )

    params = {
        "start_dt": start_dt,
        "end_dt": end_dt,
        "nivel": nivel,
        "portico_id": portico_id,
        "vehicle_id": vehicle_id,
    }

    res = await db.execute(sql, params)
    return [dict(r) for r in res.mappings().all()]


@router.get("/crew-productive-time")

async def crew_productive_time(
    start_date: str | None = None,
    end_date: str | None = None,
    nivel: int | None = None,
    portico: str | None = None,
    db: AsyncSession = Depends(get_db),
):
    """Tiempo productivo de cuadrilla por zona.

    Usa la tabla `zenit_crew_productive_time_generated` (seedeada en zenit_demo_1_dia.sql).
    """
    sql = text("""
        SELECT crew_name, zone_name, day::date as day, productive_minutes, nivel
        FROM zenit_crew_productive_time_generated
        WHERE (CAST(:start_date AS date) IS NULL OR day::date >= CAST(:start_date AS date))
            AND (CAST(:end_date AS date) IS NULL OR day::date <= CAST(:end_date AS date))
            AND (CAST(:nivel AS integer) IS NULL OR nivel = CAST(:nivel AS integer))
            AND (CAST(:portico AS text) IS NULL OR portico_code = CAST(:portico AS text))
        ORDER BY day::date ASC, crew_name ASC, zone_name ASC
    """)
    result = await db.execute(sql, {"start_date": start_date, "end_date": end_date, "nivel": nivel, "portico": portico})
    rows = result.mappings().all()

    # Front exporta y grafica usando: fecha, cuadrilla, nivel, zona, minutos
    out = []
    for r in rows:
        out.append({
            "fecha": str(r.get("day")),
            "cuadrilla": r.get("crew_name"),
            "nivel": r.get("nivel"),
            "zona": r.get("zone_name"),
            "minutos": int(r.get("productive_minutes") or 0),
        })

    return out

@router.get("/tactical/camera-availability")
async def tactical_camera_availability(
    start_date: date | None = Query(None),
    end_date: date | None = Query(None),
    nivel: int | None = None,
    db: AsyncSession = Depends(get_db),
):
    if not await _table_exists(db, "zenit_tactical_camera_availability_generated"):
        return []

    start_dt, end_dt = _build_datetime_range(start_date, end_date)
    result = await db.execute(
        text(
            """
            SELECT
                CASE nivel
                    WHEN 1 THEN 'NP'
                    WHEN 2 THEN 'UCL'
                    WHEN 3 THEN 'NTI'
                    ELSE CONCAT('Nivel ', nivel)
                END AS nivel,
                SUM(online)::int AS online,
                SUM(offline)::int AS offline,
                SUM(mantenimiento)::int AS mantenimiento
            FROM zenit_tactical_camera_availability_generated
            WHERE day >= :start_dt
              AND day < :end_dt
              AND (CAST(:nivel AS integer) IS NULL OR nivel = CAST(:nivel AS integer))
            GROUP BY 1
            ORDER BY 1;
            """
        ),
        {"start_dt": start_dt, "end_dt": end_dt, "nivel": nivel},
    )
    return result.mappings().all()

@router.get("/tactical/camera-alerts")
async def tactical_camera_alerts(
    start_date: date | None = Query(None),
    end_date: date | None = Query(None),
    nivel: int | None = None,
    db: AsyncSession = Depends(get_db),
):
    start_dt, end_dt = _build_datetime_range(start_date, end_date)
    if await _table_exists(db, "camera_finding_event"):
        result = await db.execute(
            text(
                """
                SELECT
                    COALESCE(NULLIF(finding_label, ''), finding_code) AS tipo,
                    COUNT(*)::int AS cantidad
                FROM camera_finding_event cfe
                LEFT JOIN camera_video cv
                  ON cv.id = cfe.video_id
                WHERE cfe.detected_at >= :start_dt
                  AND cfe.detected_at < :end_dt
                  AND (
                    CAST(:nivel AS integer) IS NULL
                    OR UPPER(COALESCE(cv.nivel, '')) = CASE CAST(:nivel AS integer)
                        WHEN 1 THEN 'NP'
                        WHEN 2 THEN 'UCL'
                        WHEN 3 THEN 'NTI'
                        ELSE UPPER(COALESCE(cv.nivel, ''))
                    END
                  )
                GROUP BY 1
                ORDER BY cantidad DESC, tipo ASC;
                """
            ),
            {"start_dt": start_dt, "end_dt": end_dt, "nivel": nivel},
        )
        live_rows = result.mappings().all()
        if live_rows:
            return live_rows

    if not await _table_exists(db, "zenit_tactical_camera_alerts_generated"):
        return []

    result = await db.execute(
        text(
            """
            SELECT
                tipo,
                SUM(cantidad)::int AS cantidad
            FROM zenit_tactical_camera_alerts_generated
            WHERE day >= :start_dt
              AND day < :end_dt
              AND (CAST(:nivel AS integer) IS NULL OR nivel = CAST(:nivel AS integer))
            GROUP BY tipo
            ORDER BY cantidad DESC, tipo ASC;
            """
        ),
        {"start_dt": start_dt, "end_dt": end_dt, "nivel": nivel},
    )
    return result.mappings().all()


@router.get("/tactical/camera-finding-summary")
async def tactical_camera_finding_summary(
    start_date: date | None = Query(None),
    end_date: date | None = Query(None),
    nivel: int | None = None,
    db: AsyncSession = Depends(get_db),
):
    if not await _table_exists(db, "camera_finding_snapshot"):
        return []

    start_dt, end_dt = _build_datetime_range(start_date, end_date)
    result = await db.execute(
        text(
            """
            SELECT
                cfs.finding_code AS code,
                cfs.finding_group AS grupo,
                cfs.finding_label AS hallazgo,
                cfs.coverage AS cobertura,
                SUM(cfs.active_count)::int AS activos,
                SUM(cfs.total_count)::int AS acumulados,
                MAX(cfs.last_seen_at) AS ultimo_evento,
                MAX(NULLIF(cfs.last_evidence, '')) AS evidencia
            FROM camera_finding_snapshot cfs
            LEFT JOIN camera_video cv
              ON cv.id = cfs.video_id
            WHERE cfs.updated_at >= :start_dt
              AND cfs.updated_at < :end_dt
              AND (
                CAST(:nivel AS integer) IS NULL
                OR UPPER(COALESCE(cv.nivel, '')) = CASE CAST(:nivel AS integer)
                    WHEN 1 THEN 'NP'
                    WHEN 2 THEN 'UCL'
                    WHEN 3 THEN 'NTI'
                    ELSE UPPER(COALESCE(cv.nivel, ''))
                END
              )
            GROUP BY 1, 2, 3, 4
            ORDER BY acumulados DESC, activos DESC, code ASC;
            """
        ),
        {"start_dt": start_dt, "end_dt": end_dt, "nivel": nivel},
    )
    return result.mappings().all()
