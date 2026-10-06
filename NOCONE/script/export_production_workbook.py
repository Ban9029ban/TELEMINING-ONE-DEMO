from __future__ import annotations

import argparse
import json
import math
import re
from collections import defaultdict
from copy import deepcopy
from datetime import date, datetime, time, timedelta
from hashlib import md5
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


DEFAULT_INPUT = Path(r"C:\Users\pipe_\Downloads\3.1. Control App SdC O-222 02-06-2026 TB.xlsm")
DEFAULT_OUTPUT = Path("web_front/src/data/productionWorkbookSeed.json")

DEFAULT_PERMISSIONS = {
    "allowedRoles": ["admin_produccion"],
}

MAPS = [
    {"code": "NP_dia", "label": "Diamante - NP"},
    {"code": "NTI_andes_norte", "label": "NTI Andes Norte"},
    {"code": "NTI_dia", "label": "NTI Dia"},
    {"code": "UCL_dia", "label": "UCL"},
]

LEVEL_TO_MAP = {
    "NP": "NP_dia",
    "NTI": "NTI_andes_norte",
    "UCL": "UCL_dia",
    "SNV": "NTI_dia",
}


def serialize_cell(value: Any) -> Any:
    if isinstance(value, datetime):
        return value.isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, time):
        return value.strftime("%H:%M:%S")
    if isinstance(value, timedelta):
        total_seconds = int(value.total_seconds())
        hours = total_seconds // 3600
        minutes = (total_seconds % 3600) // 60
        seconds = total_seconds % 60
        return f"{hours:02d}:{minutes:02d}:{seconds:02d}"
    return value


def slugify(value: str) -> str:
    lowered = re.sub(r"[^a-z0-9]+", "-", value.lower()).strip("-")
    return lowered or "item"


def combine_date_and_time(raw_date: Any, raw_time: Any) -> str | None:
    if not raw_date or not raw_time:
        return None

    if isinstance(raw_date, datetime):
        base_date = raw_date.date()
    elif isinstance(raw_date, date):
        base_date = raw_date
    else:
        try:
            base_date = datetime.fromisoformat(str(raw_date)).date()
        except ValueError:
            return None

    if isinstance(raw_time, time):
        parts = (raw_time.hour, raw_time.minute, raw_time.second)
    elif isinstance(raw_time, datetime):
        parts = (raw_time.hour, raw_time.minute, raw_time.second)
    else:
        text = str(raw_time).strip()
        if not text:
            return None
        try:
            parsed = datetime.strptime(text, "%H:%M:%S").time()
        except ValueError:
            try:
                parsed = datetime.strptime(text, "%H:%M").time()
            except ValueError:
                return None
        parts = (parsed.hour, parsed.minute, parsed.second)

    return datetime.combine(base_date, time(*parts)).isoformat()


def normalize_schedule_window(start_iso: str | None, end_iso: str | None) -> tuple[str | None, str | None]:
    if not start_iso or not end_iso:
        return start_iso, end_iso

    start_dt = datetime.fromisoformat(start_iso)
    end_dt = datetime.fromisoformat(end_iso)
    if end_dt < start_dt:
        end_dt += timedelta(days=1)
    return start_dt.isoformat(), end_dt.isoformat()


def normalize_duration_minutes(start_iso: str | None, end_iso: str | None) -> int:
    if not start_iso or not end_iso:
        return 0
    start_dt = datetime.fromisoformat(start_iso)
    end_dt = datetime.fromisoformat(end_iso)
    if end_dt < start_dt:
        end_dt += timedelta(days=1)
    return max(0, int((end_dt - start_dt).total_seconds() // 60))


def map_code_for_level(level: str | None) -> str:
    return LEVEL_TO_MAP.get(str(level or "").strip().upper(), "NP_dia")


def parse_control_sheet(workbook, sheet_name: str, domain: str) -> list[dict[str, Any]]:
    ws = workbook[sheet_name]
    headers = [cell for cell in next(ws.iter_rows(min_row=10, max_row=10, values_only=True))]
    tasks: list[dict[str, Any]] = []

    for row in ws.iter_rows(min_row=11, values_only=True):
        raw = {}
        for idx, header in enumerate(headers):
            if not header:
                continue
            raw[str(header)] = serialize_cell(row[idx] if idx < len(row) else None)

        if not raw.get("Fecha") or not raw.get("Postura") or not raw.get("Actividad"):
            continue

        planned_start = combine_date_and_time(raw.get("Fecha"), raw.get("Hora Inicio Prog"))
        planned_end = combine_date_and_time(raw.get("Fecha"), raw.get("Hora Término Prog"))
        real_start = combine_date_and_time(raw.get("Fecha"), raw.get("Hora Inicio Real"))
        real_end = combine_date_and_time(raw.get("Fecha"), raw.get("Hora Término Real"))
        planned_start, planned_end = normalize_schedule_window(planned_start, planned_end)
        real_start, real_end = normalize_schedule_window(real_start, real_end)

        task_id = f"{domain}-{len(tasks) + 1:04d}"
        level = str(raw.get("Nivel") or "").strip()
        posture = str(raw.get("Postura") or "").strip()
        waypoint_id = f"wp-{slugify(f'{level}-{posture}')}"
        map_code = map_code_for_level(level)

        progress = raw.get("Avance (%)")
        try:
            progress_value = float(progress)
        except (TypeError, ValueError):
            progress_value = 0.0

        planned_minutes = normalize_duration_minutes(planned_start, planned_end)
        real_minutes = normalize_duration_minutes(real_start, real_end)
        delay_minutes = max(real_minutes - planned_minutes, 0) if planned_minutes else 0

        linked_roles = []
        for key in ("Cargo 1", "Cargo 2", "Cargo 3", "Cargo 4", "Cargo 5"):
            value = str(raw.get(key) or "").strip()
            if value:
                linked_roles.append(value)

        task = {
            "id": task_id,
            "domain": domain,
            "sheetName": sheet_name,
            "date": raw.get("Fecha"),
            "turno": raw.get("Turno"),
            "nivel": level,
            "postura": posture,
            "mapCode": map_code,
            "waypointId": waypoint_id,
            "type": raw.get("Tipo"),
            "activity": raw.get("Actividad"),
            "priority": raw.get("Prioridad"),
            "progress": progress_value,
            "plannedStart": planned_start,
            "plannedEnd": planned_end,
            "plannedMinutes": planned_minutes,
            "realStart": real_start,
            "realEnd": real_end,
            "realMinutes": real_minutes,
            "delayMinutes": delay_minutes,
            "equipmentPlanned": raw.get("Equipo Prog"),
            "crewPlanned": raw.get("Cuadrilla Prog"),
            "equipmentReal": raw.get("Equipo Real"),
            "crewReal": raw.get("Cuadrilla Real"),
            "cncExternalCategory": raw.get("Categoría CNC Externa"),
            "cncExternalSubcategory": raw.get("Subcategoría CNC Externa"),
            "cncExternalDetail": raw.get("Detalle CNC Externa"),
            "cncExternalStart": combine_date_and_time(raw.get("Fecha"), raw.get("Inicio CNC Externa")),
            "cncExternalEnd": combine_date_and_time(raw.get("Fecha"), raw.get("Término CNC Externa")),
            "cncInternalCategory": raw.get("Categoría CNC Interna"),
            "cncInternalSubcategory": raw.get("Subcategoría CNC Interna"),
            "cncInternalDetail": raw.get("Detalle de CNC Interna"),
            "cncInternalStart": combine_date_and_time(raw.get("Fecha"), raw.get("Inicio CNC Interna")),
            "cncInternalEnd": combine_date_and_time(raw.get("Fecha"), raw.get("Término CNC Interna")),
            "roles": linked_roles,
            "advancedStay": raw.get("Avanzada - Quedada"),
            "observation": raw.get("Observación"),
            "hip": raw.get("HIP"),
            "htp": raw.get("HTP"),
            "hir": raw.get("HIR"),
            "htr": raw.get("HTR"),
            "hie": raw.get("HIE"),
            "hte": raw.get("HTE"),
            "hii": raw.get("HII"),
            "hti": raw.get("HTI"),
            "trackingPorticoId": None,
            "trackingPorticoName": "",
            "linkedPeople": [],
            "linkedVehicles": [],
            "status": derive_task_status(progress_value, planned_start, planned_end, real_start, real_end),
        }
        tasks.append(task)

    return tasks


def derive_task_status(
    progress: float,
    planned_start: str | None,
    planned_end: str | None,
    real_start: str | None,
    real_end: str | None,
) -> str:
    if progress >= 1:
        return "completed"
    if real_start and not real_end:
        return "in_progress"
    if planned_start and not real_start:
        return "planned"
    if planned_end and real_end:
        planned_end_dt = datetime.fromisoformat(planned_end)
        real_end_dt = datetime.fromisoformat(real_end)
        if real_end_dt > planned_end_dt:
            return "delayed"
    if progress > 0:
        return "partially_complete"
    return "planned"


def parse_gantt_sheet(workbook, sheet_name: str, domain: str) -> dict[str, Any]:
    ws = workbook[sheet_name]
    date_value = serialize_cell(ws["C1"].value)
    turno_value = serialize_cell(ws["C2"].value)
    cumplimiento_value = ws["I2"].value
    try:
        cumplimiento = float(cumplimiento_value)
    except (TypeError, ValueError):
        cumplimiento = 0.0

    slot_columns: list[tuple[int, str]] = []
    header_row = [cell for cell in next(ws.iter_rows(min_row=5, max_row=5, values_only=True))]
    for idx, value in enumerate(header_row, start=1):
        if idx < 9:
            continue
        text = str(value or "").strip()
        if text:
            slot_columns.append((idx, text))

    rows: list[dict[str, Any]] = []
    current_row = 6
    while current_row <= ws.max_row:
        number_value = ws.cell(row=current_row, column=1).value
        posture_value = ws.cell(row=current_row, column=3).value
        if not number_value and not posture_value:
            current_row += 1
            continue

        program_row = current_row
        resource_program_row = current_row + 1
        real_row = current_row + 2
        resource_real_row = current_row + 3
        level = str(ws.cell(row=program_row, column=2).value or "").strip()
        posture = str(ws.cell(row=program_row, column=3).value or "").strip()
        if not posture:
            current_row += 4
            continue

        waypoint_id = f"wp-{slugify(f'{level}-{posture}')}"
        slots = []
        for col_idx, slot_label in slot_columns:
            planned_activity = serialize_cell(ws.cell(row=program_row, column=col_idx).value)
            planned_resource = serialize_cell(ws.cell(row=resource_program_row, column=col_idx).value)
            real_activity = serialize_cell(ws.cell(row=real_row, column=col_idx).value)
            real_resource = serialize_cell(ws.cell(row=resource_real_row, column=col_idx).value)
            if any([planned_activity, planned_resource, real_activity, real_resource]):
                slots.append(
                    {
                        "slot": slot_label,
                        "plannedActivity": planned_activity,
                        "plannedResource": planned_resource,
                        "realActivity": real_activity,
                        "realResource": real_resource,
                    }
                )

        rows.append(
            {
                "id": f"{domain}-gantt-{len(rows) + 1:04d}",
                "domain": domain,
                "nivel": level,
                "postura": posture,
                "waypointId": waypoint_id,
                "mapCode": map_code_for_level(level),
                "observations": serialize_cell(ws.cell(row=program_row, column=8).value),
                "slots": slots,
            }
        )
        current_row += 4

    return {
        "sheetName": sheet_name,
        "date": date_value,
        "turno": turno_value,
        "cumplimiento": cumplimiento,
        "timeSlots": [label for _, label in slot_columns],
        "rows": rows,
    }


def parse_hours_sheet(workbook, sheet_name: str) -> dict[str, Any]:
    ws = workbook[sheet_name]

    def to_hours(raw_value: Any) -> float:
        if isinstance(raw_value, timedelta):
            return round(raw_value.total_seconds() / 3600, 2)
        if isinstance(raw_value, time):
            delta = timedelta(hours=raw_value.hour, minutes=raw_value.minute, seconds=raw_value.second)
            return round(delta.total_seconds() / 3600, 2)
        text = str(raw_value or "").strip()
        if not text:
            return 0.0
        if "day" in text:
            day_part, clock_part = text.split(",")
            days = float(day_part.split()[0])
            parsed = datetime.strptime(clock_part.strip(), "%H:%M:%S").time()
            total = timedelta(days=days, hours=parsed.hour, minutes=parsed.minute, seconds=parsed.second)
            return round(total.total_seconds() / 3600, 2)
        try:
            parsed = datetime.strptime(text, "%H:%M:%S").time()
            return round((parsed.hour * 3600 + parsed.minute * 60 + parsed.second) / 3600, 2)
        except ValueError:
            return 0.0

    crew_rows = []
    for row_idx in range(3, 22):
        label = ws.cell(row=row_idx, column=1).value
        if not label:
            continue
        crew_rows.append(
            {
                "label": str(label),
                "productiveHours": to_hours(ws.cell(row=row_idx, column=2).value),
                "targetHours": to_hours(ws.cell(row=row_idx, column=3).value),
            }
        )

    equipment_rows = []
    for row_idx in range(24, ws.max_row + 1):
        label = ws.cell(row=row_idx, column=1).value
        if not label:
            continue
        equipment_rows.append(
            {
                "label": str(label),
                "productiveHours": to_hours(ws.cell(row=row_idx, column=2).value),
            }
        )

    return {
        "sheetName": sheet_name,
        "crew": crew_rows,
        "equipment": equipment_rows,
    }


def parse_cnc_sheet(workbook, sheet_name: str) -> dict[str, Any]:
    ws = workbook[sheet_name]

    def parse_category_matrix(header_row: int, start_row: int, end_row: int, start_col: int, end_col: int) -> list[dict[str, Any]]:
        categories = []
        for col_idx in range(start_col, end_col + 1):
            category = ws.cell(row=header_row, column=col_idx).value
            if not category:
                continue
            options = []
            for row_idx in range(start_row, end_row + 1):
                option_value = ws.cell(row=row_idx, column=col_idx).value
                if option_value:
                    options.append(str(option_value))
            categories.append(
                {
                    "category": str(category),
                    "options": dedupe_preserve_order(options),
                }
            )
        return categories

    cargos_mineria = []
    for row_idx in range(45, 61):
        value = ws.cell(row=row_idx, column=2).value
        if value:
            cargos_mineria.append(str(value))

    cargos_oocc = []
    for row_idx in range(65, 85):
        value = ws.cell(row=row_idx, column=2).value
        if value:
            cargos_oocc.append(str(value))

    return {
        "external": parse_category_matrix(1, 2, 18, 4, 18),
        "internal": parse_category_matrix(21, 22, 41, 4, 16),
        "cargosMineria": cargos_mineria,
        "cargosOOCC": cargos_oocc,
    }


def parse_lists_sheet(workbook, sheet_name: str) -> dict[str, Any]:
    ws = workbook[sheet_name]
    sections: dict[str, list[dict[str, str]]] = {}
    header_row = [cell for cell in next(ws.iter_rows(min_row=1, max_row=1, values_only=True))]

    for col_idx in range(1, len(header_row) + 1, 2):
        header = header_row[col_idx - 1]
        if not header:
            continue

        items = []
        for row_idx in range(2, ws.max_row + 1):
            code = ws.cell(row=row_idx, column=col_idx).value
            label = ws.cell(row=row_idx, column=col_idx + 1).value if col_idx + 1 <= ws.max_column else None
            if code:
                items.append(
                    {
                        "code": str(code),
                        "label": str(label) if label else str(code),
                    }
                )

        sections[str(header)] = dedupe_dict_items(items, key_name="code")

    return sections


def dedupe_dict_items(items: list[dict[str, str]], key_name: str) -> list[dict[str, str]]:
    seen = set()
    output = []
    for item in items:
        key = item.get(key_name)
        if key in seen:
            continue
        seen.add(key)
        output.append(item)
    return output


def dedupe_preserve_order(items: list[str]) -> list[str]:
    seen = set()
    output = []
    for item in items:
        if item in seen:
            continue
        seen.add(item)
        output.append(item)
    return output


def waypoint_hash_offset(seed_value: str) -> tuple[float, float]:
    digest = md5(seed_value.encode("utf-8")).hexdigest()
    x_component = int(digest[:4], 16) / 65535
    y_component = int(digest[4:8], 16) / 65535
    return (x_component - 0.5) * 0.06, (y_component - 0.5) * 0.06


def generate_waypoints(tasks: list[dict[str, Any]]) -> list[dict[str, Any]]:
    grouped: dict[str, list[tuple[str, str]]] = defaultdict(list)
    meta_by_key: dict[str, dict[str, Any]] = {}

    for task in tasks:
        key = task["waypointId"]
        meta = meta_by_key.setdefault(
            key,
            {
                "id": key,
                "label": task["postura"],
                "postura": task["postura"],
                "nivel": task["nivel"],
                "mapCode": task["mapCode"],
                "domains": set(),
                "taskCounts": defaultdict(int),
                "trackingPorticoId": None,
                "trackingPorticoName": "",
            },
        )
        meta["domains"].add(task["domain"])
        meta["taskCounts"][task["domain"]] += 1
        grouped[task["mapCode"]].append((key, task["postura"]))

    positioned: dict[str, tuple[float, float]] = {}
    for map_code, entries in grouped.items():
        unique_keys = sorted({key for key, _ in entries})
        total = len(unique_keys)
        if total == 0:
            continue

        columns = min(5, max(2, math.ceil(math.sqrt(total))))
        rows = math.ceil(total / columns)
        x_span = 0.68
        y_span = 0.62
        x_start = 0.16
        y_start = 0.18

        for index, key in enumerate(unique_keys):
            col = index % columns
            row = index // columns
            x_step = x_span / max(columns - 1, 1)
            y_step = y_span / max(rows - 1, 1)
            x = x_start + col * x_step
            y = y_start + row * y_step
            offset_x, offset_y = waypoint_hash_offset(key)
            positioned[key] = (
                round(min(0.9, max(0.1, x + offset_x)), 4),
                round(min(0.9, max(0.1, y + offset_y)), 4),
            )

    result = []
    for key, meta in sorted(meta_by_key.items(), key=lambda item: (item[1]["mapCode"], item[1]["nivel"], item[1]["postura"])):
        x_pct, y_pct = positioned.get(key, (0.5, 0.5))
        result.append(
            {
                **meta,
                "domains": sorted(meta["domains"]),
                "taskCounts": {
                    "mineria": int(meta["taskCounts"].get("mineria", 0)),
                    "oocc": int(meta["taskCounts"].get("oocc", 0)),
                },
                "xPct": x_pct,
                "yPct": y_pct,
            }
        )

    return result


def build_payload(workbook_path: Path) -> dict[str, Any]:
    wb = load_workbook(workbook_path, data_only=True, read_only=False)

    control_mineria = parse_control_sheet(wb, "Control Minería", "mineria")
    control_oocc = parse_control_sheet(wb, "Control OO.CC", "oocc")
    all_tasks = control_mineria + control_oocc

    gantt_mineria = parse_gantt_sheet(wb, "Gantt Minería", "mineria")
    gantt_oocc = parse_gantt_sheet(wb, "Gantt OO.CC", "oocc")
    hours_mineria = parse_hours_sheet(wb, "Horas Minería")
    cnc = parse_cnc_sheet(wb, "CNC")
    lists = parse_lists_sheet(wb, "Listas")
    waypoints = generate_waypoints(all_tasks)

    operation_date = control_mineria[0]["date"] if control_mineria else control_oocc[0]["date"] if control_oocc else None

    return {
        "meta": {
            "sourceWorkbook": workbook_path.name,
            "operationDate": operation_date,
            "generatedAt": datetime.utcnow().isoformat() + "Z",
            "maps": deepcopy(MAPS),
        },
        "permissions": deepcopy(DEFAULT_PERMISSIONS),
        "referenceData": {
            "lists": lists,
            "cnc": cnc,
        },
        "dashboard": {
            "cumplimientoMineria": gantt_mineria["cumplimiento"],
            "cumplimientoOOCC": gantt_oocc["cumplimiento"],
        },
        "schedule": {
            "waypoints": waypoints,
            "controlMineria": control_mineria,
            "controlOOCC": control_oocc,
            "ganttMineria": gantt_mineria,
            "ganttOOCC": gantt_oocc,
            "hoursMineria": hours_mineria,
        },
    }


def main() -> None:
    parser = argparse.ArgumentParser(description="Exporta el Excel de produccion a un JSON util para el frontend.")
    parser.add_argument("--input", type=Path, default=DEFAULT_INPUT)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    args = parser.parse_args()

    payload = build_payload(args.input)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    with args.output.open("w", encoding="utf-8") as output_file:
        json.dump(payload, output_file, ensure_ascii=False, indent=2)
        output_file.write("\n")


if __name__ == "__main__":
    main()
